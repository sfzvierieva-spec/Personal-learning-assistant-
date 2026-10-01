"""Tests for ingestion_generation, with a fake LLM (no API key needed).

Run:
    pytest ingestion_generation/tests
"""

from __future__ import annotations

import io
import json
from datetime import date

import pytest
from docx import Document

from .. import generator
from ..chunking import chunk_text, estimate_tokens
from ..extraction import ExtractionError, extract_text
from ..generator import GenerationError, generate_content
from profile_and_prompts.profile.builder import build_profile
from profile_and_prompts.prompt_generator.generator import generate_system_prompt

COURSE = """[Page 1]
La photosynthèse est le processus par lequel les plantes convertissent l'énergie lumineuse en énergie chimique.

[Page 2]
Elle a lieu dans les chloroplastes et produit du glucose et de l'oxygène."""

ANSWERS = {
    "q1_main_format": ["flashcards", "quiz"],
    "q2_density": "dense",
    "q3_explanation_style": "definition_only",
    "q4_struggle": ["hard_to_remember"],
    "q5_study_goal": "mcq_exam",
    "q6_self_testing": "mainly_testing",
    "q7_tone": "raw_notes",
    "q8_visual_layout": "bullet_lists",
}
SYSTEM_PROMPT = generate_system_prompt(build_profile(ANSWERS))

FLASHCARDS_ANSWER = json.dumps({
    "title": "La photosynthèse",
    "language": "fr",
    "content": {"cards": [
        {"front": "Où a lieu la photosynthèse ?", "back": "Dans les chloroplastes.", "source_hint": "Page 2"},
        {"front": "Que produit-elle ?", "back": "Du glucose et de l'oxygène.", "source_hint": "Page 2"},
    ]},
}, ensure_ascii=False)


class FakeLLM:
    """Returns the queued answers in order and records every call."""

    def __init__(self, *answers: str):
        self.answers = list(answers)
        self.calls: list[tuple[str, list[dict]]] = []

    def __call__(self, system: str, messages: list[dict]) -> str:
        self.calls.append((system, [dict(m) for m in messages]))
        return self.answers.pop(0) if len(self.answers) > 1 else self.answers[0]


# --------------------------------------------------------------------------- #
# Extraction.
# --------------------------------------------------------------------------- #
def test_extract_txt_from_bytes():
    assert "photosynthèse" in extract_text(COURSE.encode("utf-8"), filename="cours.txt")


def test_extract_docx_with_table():
    doc = Document()
    doc.add_paragraph("Chapitre 1 : la cellule")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Mitochondrie"
    table.rows[0].cells[1].text = "Produit l'énergie"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = extract_text(buffer.getvalue(), filename="cours.docx")
    assert "Chapitre 1" in text
    assert "Mitochondrie | Produit l'énergie" in text


def test_extract_rejects_unknown_type():
    with pytest.raises(ExtractionError):
        extract_text(b"...", filename="slides.pptx")


def test_extract_rejects_empty_file():
    with pytest.raises(ExtractionError):
        extract_text(b"   \n  ", filename="vide.txt")


# --------------------------------------------------------------------------- #
# Chunking.
# --------------------------------------------------------------------------- #
def test_short_text_is_one_chunk():
    assert chunk_text(COURSE, max_tokens=1000) == [COURSE]


def test_long_text_chunks_respect_the_budget_and_keep_everything():
    paragraphs = [f"Paragraph {i}. " + "word " * 80 for i in range(100)]
    text = "\n\n".join(paragraphs)
    chunks = chunk_text(text, max_tokens=500, overlap_tokens=50)

    assert len(chunks) > 1
    assert all(estimate_tokens(c) <= 500 + 1 for c in chunks)
    for p in paragraphs:  # nothing lost
        assert any(p in c for c in chunks)


def test_giant_paragraph_is_split():
    chunks = chunk_text("x" * 10_000, max_tokens=500)
    assert len(chunks) > 1
    assert "".join(chunks) == "x" * 10_000


# --------------------------------------------------------------------------- #
# Generation (single pass).
# --------------------------------------------------------------------------- #
def test_flashcards_single_pass():
    llm = FakeLLM(FLASHCARDS_ANSWER)
    result = generate_content(COURSE, SYSTEM_PROMPT, "flashcards", {"num_items": 2}, llm=llm)

    assert result["format"] == "flashcards"
    assert result["language"] == "fr"
    assert len(result["content"]["cards"]) == 2
    assert result["meta"]["strategy"] == "single_pass"
    assert result["meta"]["warnings"] == []

    system, messages = llm.calls[0]
    assert system == SYSTEM_PROMPT                    # Léna's prompt is used as-is
    assert "<course_material>" in messages[0]["content"]
    assert "Create 2 flashcards" in messages[0]["content"]


def test_answer_in_markdown_fence_is_accepted():
    llm = FakeLLM(f"```json\n{FLASHCARDS_ANSWER}\n```")
    result = generate_content(COURSE, SYSTEM_PROMPT, "flashcards", {"num_items": 2}, llm=llm)
    assert len(result["content"]["cards"]) == 2


def test_invalid_json_triggers_one_retry():
    llm = FakeLLM("Sure! Here are your flashcards: ...", FLASHCARDS_ANSWER)
    result = generate_content(COURSE, SYSTEM_PROMPT, "flashcards", {"num_items": 2}, llm=llm)

    assert result["meta"]["retried"] is True
    retry_messages = llm.calls[1][1]
    assert [m["role"] for m in retry_messages] == ["user", "assistant", "user"]
    assert "could not be used" in retry_messages[-1]["content"]


def test_invalid_twice_raises():
    llm = FakeLLM("not json", "still not json")
    with pytest.raises(GenerationError):
        generate_content(COURSE, SYSTEM_PROMPT, "flashcards", llm=llm)


def test_quiz_with_wrong_answer_index_is_rejected():
    bad = json.dumps({"title": "t", "language": "fr", "content": {"questions": [
        {"question": "?", "choices": ["a", "b"], "answer_index": 5, "explanation": "e"}]}})
    with pytest.raises(GenerationError):
        generate_content(COURSE, SYSTEM_PROMPT, "quiz", llm=FakeLLM(bad, bad))


def test_fewer_items_than_requested_gives_a_warning():
    result = generate_content(COURSE, SYSTEM_PROMPT, "flashcards", {"num_items": 10},
                              llm=FakeLLM(FLASHCARDS_ANSWER))
    assert "2 items produced instead of the 10 requested." in result["meta"]["warnings"]


def test_synthesis_sheet_alias_maps_to_summary():
    answer = json.dumps({"title": "t", "language": "fr", "content": {
        "sections": [{"heading": "Définition", "points": ["..."]}], "key_terms": []}})
    result = generate_content(COURSE, SYSTEM_PROMPT, "synthesis_sheet", llm=FakeLLM(answer))
    assert result["format"] == "summary"


def test_unknown_format_and_empty_inputs():
    with pytest.raises(GenerationError):
        generate_content(COURSE, SYSTEM_PROMPT, "podcast", llm=FakeLLM("{}"))
    with pytest.raises(GenerationError):
        generate_content("   ", SYSTEM_PROMPT, "quiz", llm=FakeLLM("{}"))
    with pytest.raises(GenerationError):
        generate_content(COURSE, "", "quiz", llm=FakeLLM("{}"))


def test_revision_plan_length_comes_from_exam_date(monkeypatch):
    monkeypatch.setattr(generator, "_today", lambda: date(2026, 10, 1))
    answer = json.dumps({"title": "t", "language": "fr", "content": {"days": [
        {"day": i, "focus": "f", "tasks": ["t"], "duration_minutes": 30} for i in range(1, 6)]}})
    llm = FakeLLM(answer)
    result = generate_content(COURSE, SYSTEM_PROMPT, "revision_plan",
                              {"exam_date": "2026-10-06", "minutes_per_day": 30}, llm=llm)

    assert result["meta"]["user_context"]["num_items"] == 5
    assert "over exactly 5 days" in llm.calls[0][1][0]["content"]
    assert result["meta"]["warnings"] == []


# --------------------------------------------------------------------------- #
# Generation (long course -> map/reduce).
# --------------------------------------------------------------------------- #
def test_long_course_uses_map_reduce(monkeypatch):
    monkeypatch.setattr(generator, "SINGLE_PASS_MAX_TOKENS", 300)
    monkeypatch.setattr(generator, "CHUNK_TOKENS", 200)
    long_course = "\n\n".join(f"[Page {i}]\n" + "La cellule contient un noyau. " * 20
                              for i in range(1, 11))

    calls: list[tuple[str, list[dict]]] = []

    def fake(system, messages):
        calls.append((system, messages))
        if system == generator.NOTES_SYSTEM_PROMPT:
            return "- note courte (Page 1)"
        return FLASHCARDS_ANSWER

    result = generate_content(long_course, SYSTEM_PROMPT, "flashcards", {"num_items": 2}, llm=fake)

    assert result["meta"]["strategy"] == "map_reduce"
    assert result["meta"]["chunks"] > 1
    note_calls = [c for c in calls if c[0] == generator.NOTES_SYSTEM_PROMPT]
    assert len(note_calls) == result["meta"]["chunks"]
    final_system, final_messages = calls[-1]
    assert "<course_notes>" in final_messages[0]["content"]
    assert final_system == SYSTEM_PROMPT  # personalization only on the final call
