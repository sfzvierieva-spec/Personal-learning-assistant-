"""Text extraction from uploaded course files (PDF, DOCX, TXT/MD).

Public function:

- ``extract_text(source, filename=None) -> str``
      ``source`` can be a path, raw bytes, or a file-like object with ``.read()``
      (a Streamlit ``UploadedFile`` works as-is, its ``.name`` gives the type).

PDF pages are separated by ``[Page N]`` markers so the LLM can point back to
where a fact comes from (useful to check hallucinations).
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import BinaryIO, Union

from docx import Document
from pypdf import PdfReader

Source = Union[str, Path, bytes, BinaryIO]

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class ExtractionError(Exception):
    """Raised when a file cannot be read or contains no usable text."""


# --------------------------------------------------------------------------- #
# Helpers.
# --------------------------------------------------------------------------- #
def _read_source(source: Source, filename: str | None) -> tuple[bytes, str]:
    """Return (raw bytes, lowercase extension) whatever the input type."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise ExtractionError(f"File not found: {path}")
        return path.read_bytes(), path.suffix.lower()

    if isinstance(source, bytes):
        data = source
    else:
        data = source.read()
        filename = filename or getattr(source, "name", None)

    if not filename:
        raise ExtractionError("A filename is needed to know the file type.")
    return data, Path(filename).suffix.lower()


def _clean(text: str) -> str:
    """Normalize whitespace without destroying paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)   # words cut by PDF hyphenation
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# One reader per format.
# --------------------------------------------------------------------------- #
def _read_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises many different error types
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    pages = []
    for number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"[Page {number}]\n{page_text}")
    return "\n\n".join(pages)


def _read_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Tables are often used for definitions / comparisons in course notes.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _read_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".txt": _read_txt, ".md": _read_txt}


# --------------------------------------------------------------------------- #
# Public function.
# --------------------------------------------------------------------------- #
def extract_text(source: Source, filename: str | None = None) -> str:
    """Extract clean text from a course file."""
    data, extension = _read_source(source, filename)
    if extension not in _READERS:
        raise ExtractionError(
            f"Unsupported file type '{extension}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    text = _clean(_READERS[extension](data))
    if not text:
        hint = " (it may be a scanned PDF with images only)" if extension == ".pdf" else ""
        raise ExtractionError(f"No text could be extracted from the file{hint}.")
    return text
